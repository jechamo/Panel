from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter

from app.core.settings import get_settings
from app.main import app
from app.parsers.docx import parse_docx
from app.parsers.pdf import parse_pdf
from app.parsers.xlsx import parse_xlsx


def test_upload_file_stores_attachment_and_metadata(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.setenv('UPLOADS_STORAGE_DIR', str(tmp_path))
    get_settings.cache_clear()

    client = TestClient(app)
    upload = build_docx_bytes('Hello from docx')

    response = client.post(
        '/files/upload',
        data={'flow_id': 'flow-123'},
        files={
            'file': (
                'demo.docx',
                upload,
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            ),
        },
    )

    assert response.status_code == 201
    payload = response.json()['data']
    stored_dir = tmp_path / 'flow-123' / payload['id']

    assert payload['name'] == 'demo.docx'
    assert payload['variableName'] == 'demo'
    assert (stored_dir / 'source.docx').exists()
    assert (stored_dir / 'parsed.txt').read_text(encoding='utf-8') == 'Hello from docx'
    assert (stored_dir / 'metadata.json').exists()

    get_settings.cache_clear()


def test_parse_docx_returns_text() -> None:
    assert parse_docx(build_docx_bytes('Paragraph one')) == 'Paragraph one'


def test_parse_xlsx_returns_sheet_text() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Datos'
    worksheet.append(['Nombre', 'Valor'])
    worksheet.append(['alpha', '1'])
    buffer = BytesIO()
    workbook.save(buffer)

    parsed = parse_xlsx(buffer.getvalue())

    assert '[Datos]' in parsed
    assert 'Nombre | Valor' in parsed
    assert 'alpha | 1' in parsed


def test_parse_pdf_handles_blank_pdf() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)

    assert parse_pdf(buffer.getvalue()) == ''


def build_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()